from django.shortcuts import render, redirect
from .models import Transporte, Escuela, Oferente, Chofer
from .forms import TransporteForm, EscuelaForm, OferenteForm, ChoferForm
from django.db.models import Q, Sum
from openpyxl import Workbook
from io import BytesIO
from django.http import HttpResponse
# @login_required


def transporte_list(request):
    query = request.GET.get('q', '')
    if query:
        transportes = Transporte.objects.filter(
            Q(patente__icontains=query) |
            Q(oferente__nombre__icontains=query) |
            Q(chofer__nombre__icontains=query) |
            Q(sectores__icontains=query) |
            Q(escuela__nombre__icontains=query)
        )
    else:
        transportes = Transporte.objects.all()
    total_km = transportes.aggregate(Sum('cantidad_km'))['cantidad_km__sum']
    total_alumnos = transportes.aggregate(Sum('alumnos'))['alumnos__sum']
    return render(request, 'transporte/transporte_list.html', {'transportes': transportes, 'total_km': total_km, 'total_alumnos': total_alumnos})


# @login_required
def transporte_update(request, patente):
    try:
        transporte = Transporte.objects.get(patente=patente)
    except Transporte.DoesNotExist:
        return redirect('transporte_list')

    if request.method == 'POST':
        form = TransporteForm(request.POST, instance=transporte)
        if form.is_valid():
            form.save()
            form.save_m2m()
            return redirect('transporte_list')
    else:
        form = TransporteForm(instance=transporte)
    return render(request, 'transporte/transporte_form.html', {'form': form})


# @login_required
def transporte_delete(request, patente):
    try:
        transporte = Transporte.objects.get(patente=patente)
    except Transporte.DoesNotExist:
        return redirect('transporte_list')

    if request.method == 'POST':
        transporte.delete()
        return redirect('transporte_list')
    else:
        return render(request, 'transporte/transporte_confirm_delete.html', {'object': transporte})


# @login_required
def transporte_create(request):
    if request.method == 'POST':
        form = TransporteForm(request.POST)
        if form.is_valid():
            transporte = form.save(commit=False)
            escuelas = request.POST.getlist('escuela_set')
            for escuela_id in escuelas:
                escuela = Escuela.objects.get(pk=escuela_id)
                transporte.escuela.add(escuela)
            transporte.save()
            return redirect('transporte_list')
    else:
        form = TransporteForm()
    return render(request, 'transporte/transporte_form.html', {'form': form})


# @login_required
def escuela_create(request):
    if request.method == 'POST':
        form = EscuelaForm(request.POST)
        if form.is_valid():
            form.save()
            return redirect('transporte_list')
    else:
        form = EscuelaForm()
    return render(request, 'transporte/escuela_form.html', {'form': form})


def oferente_create(request):
    if request.method == 'POST':
        form = OferenteForm(request.POST)
        if form.is_valid():
            form.save()
            return redirect('transporte_list')
    else:
        form = OferenteForm()
    return render(request, 'transporte/oferente_form.html', {'form': form})


def oferente_update(request, id):
    try:
        oferente = Oferente.objects.get(id=id)
    except Oferente.DoesNotExist:
        return redirect('transporte_list')

    if request.method == 'POST':
        form = OferenteForm(instance=oferente)
        if form.is_valid():
            form.save()
            return redirect('transporte_list')
    else:
        form = OferenteForm(instance=oferente)
    return render(request, 'transporte/oferente_form.html', {'form': form})

# @login_required


def chofer_create(request):
    if request.method == 'POST':
        form = ChoferForm(request.POST)
        if form.is_valid():
            form.save()
            return redirect('transporte_list')
    else:
        form = ChoferForm()
    return render(request, 'transporte/chofer_form.html', {'form': form})

# @login_required


def chofer_update(request, id):
    try:
        chofer = Chofer.objects.get(id=id)
    except Chofer.DoesNotExist:
        return redirect('transporte_list')

    if request.method == 'POST':
        form = ChoferForm(request.POST, instance=chofer)
        if form.is_valid():
            form.save()
            return redirect('transporte_list')
    else:
        form = ChoferForm(instance=chofer)
    return render(request, 'transporte/chofer_form.html', {'form': form})


def export_data(request):
    wb = Workbook()
    ws = wb.active
    ws.title = 'Transporte Data'

    # Column names
    columns = [
        'Patente', 'Oferente', 'Rut Oferente', 'Chofer', 'Rut Chofer',
        'Cantidad de KM', 'Alumnos', 'Sectores', 'Escuela RBD', 'Escuela DV', 'URL Mapa'
    ]
    for index, column in enumerate(columns, start=1):
        ws.cell(row=1, column=index, value=column)

    # Data
    for row_num, transporte in enumerate(Transporte.objects.all(), start=2):
        escuelas = transporte.escuela.all()
        for escuela in escuelas:
            data = [
                transporte.patente,
                transporte.oferente.nombre,
                transporte.oferente.rut,
                transporte.chofer.nombre,
                transporte.chofer.rut,
                transporte.cantidad_km,
                transporte.alumnos,
                transporte.sectores,
                escuela.rbd,
                escuela.digito_verificador,
                transporte.url_mapa
            ]
            for col_num, cell_value in enumerate(data, start=1):
                ws.cell(row=row_num, column=col_num, value=cell_value)

    # Create a BytesIO object and save the workbook to it
    virtual_workbook = BytesIO()
    wb.save(virtual_workbook)

    # Prepare the response
    response = HttpResponse(content=virtual_workbook.getvalue(
    ), content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = 'attachment; filename=transporte_data.xlsx'
    return response


def profile_view(request):
    return render(request, 'transporte_list.html')
